#!/usr/bin/env python3

import csv
import sys
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.shared import Pt, RGBColor


SLEWS = (10.0, 100.0, 1000.0)
LOADS = (0.5, 10.0, 100.0)


def remove_element(element):
    element.getparent().remove(element)


def set_text(cell, value):
    cell.text = value
    for paragraph in cell.paragraphs:
        paragraph.alignment = 1


def number(value):
    return f"{value:.6g}"


def read_results(path):
    records = []
    if not path.exists():
        return records
    with path.open(newline="", encoding="utf-8") as stream:
        for row in csv.DictReader(stream):
            row["value"] = float(row["value"])
            row["slew_ps"] = float(row["slew_ps"]) if row["slew_ps"] else None
            row["load_fF"] = float(row["load_fF"]) if row["load_fF"] else None
            records.append(row)
    return records


def find(records, metric, condition=None, slew=None, load=None):
    for row in records:
        if row["metric"] != metric:
            continue
        if condition is not None and row["condition"] != condition:
            continue
        if slew is not None and row["slew_ps"] != slew:
            continue
        if load is not None and row["load_fF"] != load:
            continue
        return row["value"]
    return None


def fill_matrix(table, records, metric, scale):
    for row_index, load in enumerate(LOADS, start=1):
        for column_index, slew in enumerate(SLEWS, start=1):
            value = find(records, metric, slew=slew, load=load)
            set_text(table.cell(row_index, column_index), "Pending" if value is None else number(value * scale))


def main():
    if len(sys.argv) != 3:
        raise SystemExit("usage: populate_inv_docx.py RESULTS.csv OUTPUT.docx")

    results_path = Path(sys.argv[1]).resolve()
    output_path = Path(sys.argv[2]).resolve()
    repo_root = Path(__file__).resolve().parents[1]
    template_path = repo_root / "Combinational_Lib.docx"
    records = read_results(results_path)

    document = Document(template_path)
    tables = list(document.tables)

    if "Title" not in [style.name for style in document.styles]:
        title_style = document.styles.add_style("Title", WD_STYLE_TYPE.PARAGRAPH)
        title_style.font.name = "Calibri"
        title_style.font.size = Pt(20)
        title_style.font.bold = True
        title_style.font.color.rgb = RGBColor(0, 0, 0)
    document.paragraphs[0].text = "Inverter Post Layout Characterization"
    document.paragraphs[0].style = document.styles["Title"]
    document.paragraphs[1].text = "SKY130A TT corner at 1.8 V and 27 °C using the capacitance-extracted invX1 layout."

    for paragraph in list(document.paragraphs):
        text = paragraph.text.strip()
        if text.startswith("Related pin A:"):
            paragraph.text = "Related pin vin"
        elif text.startswith("Related pin B:") or text.startswith("Related pin C:"):
            remove_element(paragraph._element)
        elif text.startswith("Static Power"):
            paragraph.paragraph_format.page_break_before = True

    for index in (2, 3, 5, 6, 8, 9, 11, 12, 15, 16, 18, 19):
        remove_element(tables[index]._element)

    cap_table = tables[0]
    set_text(cap_table.cell(1, 0), "vin")
    while len(cap_table.rows) > 2:
        remove_element(cap_table.rows[-1]._tr)

    cap_values = {
        "rise": find(records, "input_cap", condition="rise"),
        "fall": find(records, "input_cap", condition="fall"),
        "average": find(records, "input_cap", condition="average"),
    }
    for column, condition in enumerate(("rise", "fall", "average"), start=1):
        value = cap_values[condition]
        set_text(cap_table.cell(1, column), "Pending" if value is None else number(value / 1000.0))

    fill_matrix(tables[1], records, "rise_transition", 1.0 / 1000.0)
    fill_matrix(tables[4], records, "fall_transition", 1.0 / 1000.0)
    fill_matrix(tables[7], records, "cell_rise_delay", 1.0 / 1000.0)
    fill_matrix(tables[10], records, "cell_fall_delay", 1.0 / 1000.0)
    fill_matrix(tables[14], records, "rise_dynamic_power", 1.0)
    fill_matrix(tables[17], records, "fall_dynamic_power", 1.0)

    static_table = tables[13]
    while len(static_table.rows) > 3:
        remove_element(static_table.rows[-1]._tr)
    set_text(static_table.cell(0, 0), "Condition")
    set_text(static_table.cell(1, 0), "vin = 0")
    set_text(static_table.cell(2, 0), "vin = 1")
    p0 = find(records, "static_power", condition="vin_0")
    p1 = find(records, "static_power", condition="vin_1")
    set_text(static_table.cell(1, 1), "Pending" if p0 is None else number(p0))
    set_text(static_table.cell(2, 1), "Pending" if p1 is None else number(p1))

    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)


if __name__ == "__main__":
    main()
